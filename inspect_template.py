from docx import Document
p=r"C:\Users\Dsx88\Documents\xwechat_files\wxid_wqcj84d3ej1m22_d18a\msg\file\2026-09\申报书模板  开放赛道9.3.docx"
d=Document(p)
print('P');
for i,x in enumerate(d.paragraphs):
 if x.text.strip(): print(i,repr(x.text))
print('TABLES',len(d.tables))
for i,t in enumerate(d.tables):
 print('TABLE',i)
 for row in t.rows: print(' | '.join(c.text.replace('\n',' / ') for c in row.cells))
