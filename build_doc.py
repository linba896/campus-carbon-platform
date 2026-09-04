from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

out='校园能碳平台建设思路与落地方案.docx'
doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=doc.styles
styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(45,55,65)
for s,size,color in [('Title',28,'0B5D58'),('Heading 1',17,'0B5D58'),('Heading 2',13,'168B82'),('Heading 3',11,'3A6B68')]:
 st=styles[s]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6)

def shade(cell, fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell(cell,text,bold=False,color='2D3741'):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(text); r.bold=bold; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(9.5); r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
 for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],'0B5D58')
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row): set_cell(cells[i],str(v)); shade(cells[i],'F1F7F6' if len(t.rows)%2==0 else 'FFFFFF')
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 return t

def callout(label,text):
 t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; c=t.cell(0,0); shade(c,'E5F3F0'); c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string('0B5D58'); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r2=p.add_run(text); r2.font.name='Microsoft YaHei'; r2._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r2.font.size=Pt(10)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(35); p.paragraph_format.space_after=Pt(8); r=p.add_run('校园能碳平台'); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string('0B5D58'); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('建设思路与落地方案'); r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string('168B82'); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18); r=p.add_run('从数据采集、碳核算到可视化与算法迭代的完整链路'); r.italic=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string('66757F')
callout('核心判断','平台的关键不是一次性做出复杂模型，而是建立“数据可获取、公式可配置、结果可解释、系统可迭代”的轻量化闭环。')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(35); r=p.add_run('方案稿｜2026年9月'); r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string('7A8890')
doc.add_page_break()

h=doc.add_heading('一、项目定位与总体思路',1)
doc.add_paragraph('校园能碳平台面向学校后勤、节能管理与信息化部门，围绕“能耗数据—碳排放核算—空间化展示—持续优化”四个环节，形成可落地、可扩展的校园碳管理工具。首期以电力、燃气等高可得数据为主，先跑通最小可用闭环，再逐步扩展交通、用水、餐厨垃圾等数据。')
callout('总体链路','数据怎么来 → 怎么算 → 怎么展示 → 算法如何持续迭代')
doc.add_heading('建设目标',2)
for txt in ['建立建筑/分区级能耗数据台账，明确数据来源、责任部门与更新频率。','以可配置参数驱动碳排放核算，支持因子替换、版本留痕和结果复算。','用简洁的 Web/ECharts 页面呈现趋势、结构、排名和高碳热点。','为后续 ArcGIS 空间热力图、节能策略评估和校园碳盘查打基础。']:
 p=doc.add_paragraph(style='List Bullet'); p.add_run(txt)

doc.add_heading('二、数据获取：活动数据与排放因子',1)
doc.add_paragraph('校园碳排放计算需要两类基础数据：活动数据（实际消耗了多少电、气、油等）与碳排放因子（每单位消耗对应的碳排放量）。数据获取应按“优先复用既有系统、再补充调查估算”的原则推进。')
doc.add_heading('2.1 活动数据来源与责任分工',2)
table(['数据类型','建议获取方式','主要负责部门'],[
['建筑用电量','后勤/信息中心能耗监控平台；优先获取楼栋级、月度数据','后勤保障部、水电与节能管理中心'],['采暖燃气/集中供暖','财务缴费记录、燃气表或供热结算数据','后勤保障部、财务中心'],['公车油耗','车辆管理台账、加油记录','保卫处/后勤'],['水耗、食物消耗','水务/食堂采购与运营记录','后勤保障部'],['师生通勤','问卷调查；样本覆盖性别、年级、教职工等群体','项目组自行调研']], [1.4,3.0,2.1])
callout('实操建议','不必从安装新电表开始。优先联系学校后勤处或水电与节能管理中心，争取获取建筑级或分项（月度）数据；即使只有总表，也可先建立基线。')
doc.add_heading('2.2 碳排放因子来源',2)
table(['来源','特点','适用场景'],[['IPCC国家温室气体清单指南','国际通用、因子库完整','天然气、汽油、柴油等化石能源'],['国家发改委/区域电网基准线因子','中国本土化，贴合电网实际','电力排放计算'],['CLCD中国生命周期数据库','覆盖能源、材料、运输','综合碳足迹核算'],['天工LCA数据库','开放免费，数据量大、持续更新','前期建模与敏感性分析']], [2.2,2.6,1.7])
callout('关键约束','电力排放因子应优先采用中国区域电网因子（如华北、华东等），并记录因子年份与来源，避免使用全球平均值造成偏差。')

doc.add_heading('2.3 数据精度不足时的处理',2)
doc.add_paragraph('当无法获得楼栋级或分项数据时，可采用“直接计量 + 间接推算 + 估算标注”的组合方式。例如，用总电量减去已直接计量的空调、动力用电，推算照明插座用电；对餐厨垃圾等难获取数据，用估算值配合问卷补充，并在平台中标注数据质量等级与置信区间。')

doc.add_heading('三、核心算法：公式简单，参数可配置',1)
doc.add_paragraph('算法本身不追求复杂，重点是把公式、因子和边界条件做成可管理的参数。这样既便于解释，也便于后续政策口径或数据更新时快速复算。')
doc.add_heading('3.1 基础核算公式',2)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('E_total = Σ (AD_i × EF_i)'); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string('0B5D58')
doc.add_paragraph('其中：AD_i 为第 i 类能源消耗量（kWh、吨、立方米等）；EF_i 为对应的碳排放因子（如 kgCO₂/kWh）。')
doc.add_heading('3.2 分区碳排放强度',2)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('CEI_zone = (E_zone × EF_elec + Gas_zone × EF_gas) / Area_zone'); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string('0B5D58')
doc.add_paragraph('按教学楼、宿舍、食堂等分区计算单位面积碳排放强度，可横向对比并识别高碳热点。必要时可进一步引入人数、使用时长、建筑功能等修正变量。')
doc.add_heading('3.3 算法可修改的产品设计',2)
for txt in ['参数设置页：电力/燃气排放因子、建筑面积、人数、使用时长等均可编辑。','因子版本管理：保留 IPCC、中国区域电网、CLCD 等多套因子，支持切换与结果对比。','独立算法函数：将公式封装为独立 JavaScript/Python 函数，前端展示与计算逻辑解耦。','结果可追溯：每次计算记录参数版本、数据时间范围、数据质量等级和更新时间。']:
 doc.add_paragraph(txt,style='List Bullet')

doc.add_heading('四、展示方案：先轻量 Web，再空间化升级',1)
doc.add_heading('4.1 首期：ECharts + HTML 轻量页面',2)
doc.add_paragraph('首期不必引入复杂 GIS 平台，使用 ECharts 与 HTML 即可实现简洁、可解释的核心页面：')
table(['组件','展示内容','使用价值'],[['柱状图','各建筑/分区碳排放对比','快速定位高排放对象'],['饼图','电力、燃气、交通等来源结构','理解排放构成'],['折线图','月度/季度碳排放趋势','观察季节性与节能成效'],['指标卡','总排放、单位面积排放、同比变化','提供一眼可读的管理指标']], [1.2,3.0,2.3])
doc.add_heading('4.2 进阶：ArcGIS 校园碳排放热力图',2)
doc.add_paragraph('ArcGIS 的价值在于把碳排放“放回校园空间”：将建筑转为矢量多边形，结合楼层数计算建筑面积，再叠加能耗和排放因子，最终以颜色渐变表达排放强度。类似 KAUST 的实践表明，该方法适合做校园建筑级碳热点识别；USC 也用 ArcGIS Pro 制作建筑属性地图，用于隐含碳分析。')
callout('分阶段建议','第一步用 ECharts 快速跑通算法演示和管理看板；第二步在数据稳定后导入 ArcGIS Pro，制作校园碳排放空间热力图作为进阶展示。')

doc.add_heading('五、Codex 与 Kimi 的分工建议',1)
table(['工具','适合承担的任务'],[['Codex','生成 HTML + ECharts 前端页面；实现输入框、计算按钮、图表渲染；将公式封装为可配置 JavaScript/Python 函数。'],['Kimi','梳理排放因子具体数值与来源；生成模拟数据集；润色 Word 文档和技术表述。']], [1.3,5.2])
doc.add_heading('六、落地路线与首期交付物',1)
steps=[('1','数据获取','联系后勤处/水电与节能管理中心，拿到楼栋能耗数据；同步建立因子来源清单。'),('2','算法实现','用参数化公式计算总排放、分区强度和趋势指标；保留参数版本。'),('3','可视化','先交付 ECharts 看板，包含指标卡、柱状图、饼图和折线图。'),('4','验证迭代','用历史月份或模拟数据校验结果；标注估算项与置信区间。'),('5','空间升级','数据稳定后导入 ArcGIS Pro，形成校园碳排放热力图。')]
t=doc.add_table(rows=1,cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['阶段','工作重点','首期交付物']): set_cell(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],'0B5D58')
for a,b,c in steps:
 cells=t.add_row().cells; set_cell(cells[0],a,True,'0B5D58'); set_cell(cells[1],b); set_cell(cells[2],c)

doc.add_heading('七、结语：以“小闭环”换取持续迭代',1)
doc.add_paragraph('校园能碳平台的落地重点不是一次性覆盖所有排放源，而是先用最容易获得、最能反映校园运行特征的电力和燃气数据，建立一个可验证的最小闭环。随着数据质量提升，再逐步纳入交通、用水、食物与空间属性，最终形成面向节能管理、碳盘查和决策支持的校园碳管理基础设施。')
callout('下一步行动','本周优先完成两件事：① 与后勤/节能中心确认可提供的数据字段与时间粒度；② 依据确认的数据字段，制作第一版参数化 ECharts 原型。')

# footer
for section in doc.sections:
 footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=footer.add_run('校园能碳平台建设思路与落地方案'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string('8A969C'); rr.font.name='Microsoft YaHei'; rr._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.save(out)
print(out)
