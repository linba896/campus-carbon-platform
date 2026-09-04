from docx import Document
src=r'C:\Users\Dsx88\Documents\xwechat_files\wxid_wqcj84d3ej1m22_d18a\msg\file\2026-09\申报书模板  开放赛道9.3.docx'
out=r'C:\Users\Dsx88\Documents\New project\submission_feasibility_calibrated_final.docx'
d=Document(src)
repls={
'基于 AI 校园能碳算法驱动':'校园能碳算法驱动',
'运用机器学习构建 AI 校园能碳算法':'采用排放因子法、建筑类型参数和异常阈值规则构建校园能碳算法，并预留机器学习扩展接口',
'AI高碳异常识别准确率可达95%以上':'通过规则阈值识别高碳异常，识别效果将在真实标注数据充分后验证',
'校园碳排放测算误差降低18%以上':'已在模拟数据和小样本中完成核算流程验证，误差改善比例将在获得真实基准数据后评估',
'AI能碳耦合算法模型训练与优化':'校园能碳耦合算法参数校准与规则验证',
'采用监督学习算法':'采用排放因子法、建筑分类参数与规则阈值',
'模型训练数据集':'算法测试数据集（标注调研、公开或模拟数据性质）',
'AI算法模型与仿真结果材料':'校园能碳算法说明、参数表与仿真结果材料',
'异常预警准确率测试报告':'高碳预警规则测试记录',
'AI仿真模型':'参数化节能仿真模型',
'AI智能核算':'参数化能碳核算',
}
seen=set()
def change(p):
    old=p.text
    new=old
    for a,b in repls.items(): new=new.replace(a,b)
    if new!=old:
        p.text=new
        seen.add(old[:50])
for p in d.paragraphs: change(p)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: change(p)
d.add_page_break()
d.add_paragraph('附录：可实现性校准说明')
d.add_paragraph('当前可确认成果：浏览器端轻量化平台、排放因子法能碳耦合计算、参数配置、楼宇空间展示、规则预警、节能方案参数化仿真、数据导入与报告输出。机器学习训练模型、精度提升比例和异常识别准确率不作为现阶段既成成果。')
d.add_paragraph('待团队补充：AI 使用披露、参考文献、真实校园能耗样本、排放因子来源与年份、算法对照实验、异常样本人工标注、减排效益基线与成本假设。')
d.save(out)
print(out)
