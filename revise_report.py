from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src=r'C:\Users\Dsx88\Desktop\校园碳排放.docx'
out=r'C:\Users\Dsx88\Desktop\校园碳排放_平台适配修改版.docx'
d=Document(src)

repls={
'校园能碳平台面向学校后勤、节能管理与信息化部门，围绕“能耗数据—碳排放核算—空间化展示—持续优化”四个环节，形成可落地、可扩展的校园碳管理工具。首期以电力、燃气等高可得数据为主，先跑通最小可用闭环，再逐步扩展交通、用水、餐厨垃圾等数据。':'本项目已完成校园能碳可视化管控平台第一版 Web 原型。平台面向高校后勤、节能管理与信息化部门，围绕“数据导入—能碳耦合核算—空间展示—异常预警—方案仿真—报告输出”形成可运行闭环。现阶段以楼宇电力、燃气和建筑面积数据为核心，支持调研数据、模拟数据及标准化文档导入，后续可扩展交通、用水、餐厨垃圾等排放源。',
'用简洁的 Web/ECharts 页面呈现趋势、结构、排名和高碳热点。':'通过轻量化 Web 平台呈现总览指标、楼宇空间分布、趋势、来源结构、排行和高碳预警。',
'为后续 ArcGIS 空间热力图、节能策略评估和校园碳盘查打基础。':'支持上传学校校园总平面图并配置楼宇坐标，为真实校园地图适配、节能策略评估和校园碳盘查建立数据基础。',
'首期不必引入复杂 GIS 平台，使用 ECharts 与 HTML 即可实现简洁、可解释的核心页面：':'平台当前采用 React + Vite 构建轻量化 Web 原型，通过代码原生图表和交互式校园分布图实现简洁、可解释的核心页面：',
'第一步用 ECharts 快速跑通算法演示和管理看板；第二步在数据稳定后导入 ArcGIS Pro，制作校园碳排放空间热力图作为进阶展示。':'第一阶段已完成管理驾驶舱、能碳算法联动和可编辑空间分布图；第二阶段将使用真实校园总平面图校准建筑位置，数据稳定后可进一步接入 ArcGIS Pro。',
'本周优先完成两件事：① 与后勤/节能中心确认可提供的数据字段与时间粒度；② 依据确认的数据字段，制作第一版参数化 ECharts 原型。':'下一阶段优先完成三件事：① 获取真实校园总平面图并校准建筑坐标；② 按平台模板整理首批楼栋月度能耗数据；③ 用实际数据验证排放因子、预警阈值和节能仿真假设。'
}
for p in d.paragraphs:
    if p.text in repls:
        p.text=repls[p.text]

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:fill'),fill); pr.append(s)
def cell(cell,text,bold=False,color='2D4A43'):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): cell(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],'0B5D58')
    for n,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row): cell(cs[i],v); shade(cs[i],'F0F7F4' if n%2==0 else 'FFFFFF')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return t

d.add_page_break()
d.add_heading('九、平台原型现状与功能架构',1)
d.add_paragraph('当前平台名称为“源碳智算——校园能碳管控平台”，已形成可在浏览器运行的第一版 MVP。平台不依赖大规模硬件部署，可直接使用校园调研数据、模拟数据或整理后的能耗报告开展算法演示，完整对应“数字化赋能精准降碳”的参赛方向。')
table(['功能模块','当前实现','与智能控碳的关系'],[
['总览驾驶舱','综合能耗、碳排放、单位面积强度、高碳预警、趋势与来源结构','形成校园能碳运行态势总览'],
['空间能碳地图','以颜色区分楼宇排放等级，点击查看用电、排放、强度和环比','快速定位高能耗、高排放点位'],
['数据与地图中心','手工录入；导入 Word、CSV、TSV、JSON；上传校园地图；编辑楼宇坐标','降低数据归集成本并适配不同学校'],
['算法参数中心','可修改电力、天然气排放因子，结果实时联动复算','保证算法可配置、可解释、可迭代'],
['节能方案仿真','调整空调温度、照明改造覆盖率、实验室错峰比例','量化预计节电量、减排量与经济收益'],
['报告输出','生成楼宇明细、核算参数与总排放报告，可打印或另存 PDF','形成可用于汇报和方案比较的成果']],[1.35,3.15,2.15])

d.add_heading('十、数据录入与文档导入规范',1)
d.add_paragraph('平台提供两种数据输入方式。少量数据可在“数据与地图中心”直接编辑；批量数据建议使用 CSV、TSV、JSON 或包含标准表格的 Word 文档导入。导入后，数据直接进入能碳耦合模型并更新驾驶舱、空间分布和楼宇排行。')
table(['字段名称','单位/格式','用途','是否必填'],[
['建筑名称','文本','楼宇识别与地图标签','是'],['建筑类型','教学楼/实验楼/宿舍/食堂等','分类比较与场景参数选择','是'],['用电量','kWh','计算电力间接碳排放','是'],['燃气量','m³','计算燃气直接碳排放','否'],['建筑面积','m²','计算单位面积碳排放强度','是'],['环比变化','%','趋势判断与异常提示','否'],['横坐标','0—100%','建筑在学校地图中的水平位置','否'],['纵坐标','0—100%','建筑在学校地图中的垂直位置','否']],[1.35,1.35,3.1,.85])
d.add_paragraph('推荐表头：建筑名称｜建筑类型｜用电量(kWh)｜燃气量(m³)｜面积(m²)｜环比(%)｜横坐标(%)｜纵坐标(%)。Word 报告若只有叙述性文字、没有结构化楼宇能耗表格，平台将提取原文供核对，但需要在编辑表格中补充数值。')

d.add_heading('十一、学校地图适配与空间分布修改',1)
d.add_paragraph('为适配不同高校，平台允许上传 PNG 或 JPG 格式的校园总平面图作为底图。楼宇点位不写死在程序中，而是由横坐标和纵坐标两个百分比字段控制。用户可修改建筑名称、类型、面积和坐标，使空间分布图与实际学校布局对应。')
for txt in ['准备一张方向明确、建筑边界清晰的校园总平面图；','在“数据与地图中心”上传地图底图；','根据地图位置填写或调整每栋建筑的横、纵坐标；','返回总览驾驶舱，检查楼宇标签与真实位置是否一致；','后续数据条件成熟时，可将点位升级为 ArcGIS 建筑矢量多边形。']:
    d.add_paragraph(txt,style='List Number')

d.add_heading('十二、平台操作闭环与演示路线',1)
table(['步骤','操作','平台结果'],[
['1. 数据准备','整理楼栋月度用电、燃气、面积与建筑类型','形成标准数据表'],['2. 导入校验','上传文档或直接编辑数据，检查缺失项与单位','建立统一能耗台账'],['3. 参数配置','确认区域电网及燃气排放因子','确定本次核算口径'],['4. 算法核算','运行 E = Σ(AD × EF) 及单位面积强度计算','获得楼宇级碳排放结果'],['5. 地图定位','上传校园图并调整楼宇坐标','生成校园能碳空间分布'],['6. 预警仿真','识别高碳点位并调整节能策略参数','量化节电、减排与经济收益'],['7. 报告输出','点击导出能碳报告并打印/另存 PDF','形成参赛展示和管理报告']],[.9,2.85,3.0])

d.add_heading('十三、当前边界与下一阶段计划',1)
d.add_paragraph('当前版本属于前端可交互原型，重点验证校园能碳算法、数据意识、空间展示与方案仿真逻辑。数据暂存在浏览器运行状态，Word 自动识别适合结构化表格，叙述性报告仍需人工校验。下一阶段将接入持久化数据库、用户登录、历史数据版本、拖拽式地图点位编辑、正式图表库和可下载的 Word/PDF 报告服务。')

for section in d.sections:
    fp=section.footer.paragraphs[0]
    if '平台适配修改版' not in fp.text:
        fp.add_run('｜平台适配修改版')
d.save(out)
print(out)
